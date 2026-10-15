"""管口配对法兰料单回归：隔离界面导入，直接执行生产插行/清理函数。"""
import ast
from copy import copy
from io import BytesIO
from pathlib import Path
import re
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock

import openpyxl

from modules.wenbenshengcheng.nozzle_flange_materials import (
    BOLT, GASKET, NUT, PAIRED_FLANGE, STUD, build_nozzle_flange_materials,
)
from modules.wenbenshengcheng.nozzle_flange_fasteners import resolve_nozzle_fastener_dimensions


ROOT = Path(__file__).resolve().parents[1]


def attachment(code="N1", tab="配对1", tab_id="1", groups=None):
    rows = [{"参数名称": "管口号", "参数数值": code, "标题分组": None}]
    for group, params in (groups or {}).items():
        rows.extend({"标题分组": group, "参数名称": name, "参数数值": value}
                    for name, value in params.items())
    for row in rows:
        row.update({"附件类型": PAIRED_FLANGE, "Tab分类": tab, "Tab_ID": tab_id})
    return rows


def nozzle(code="N1", name="介质入口", size="100", pressure="300"):
    return {"管口代号": code, "管口功能": name, "公称尺寸": size,
            "压力等级": pressure, "法兰标准": "HG/T 20615-2009",
            "法兰型式": "WN", "密封面型式": "RF", "焊端规格": "程序推荐"}


def worksheet_namespace():
    """加载实际源码函数和模板常量，不启动产品管理器或图形界面。"""
    namespace = {"re": re, "copy": copy, "build_nozzle_flange_materials": build_nozzle_flange_materials,
                 "PAIRED_FLANGE": PAIRED_FLANGE, "db_config1": {},
                 "load_nozzle_fastener_dimensions": MagicMock(return_value={})}
    tree = ast.parse((ROOT / "modules/wenbenshengcheng/generate_material_list.py").read_text(encoding="utf-8-sig"))
    nodes = [node for node in tree.body if isinstance(node, ast.FunctionDef)
             or isinstance(node, ast.Assign) and any(
                 isinstance(target, ast.Name) and (target.id.startswith("MATERIAL_")
                                                   or target.id == "_CELL_REFERENCE_RE")
                 for target in node.targets)]
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(ROOT / "modules/wenbenshengcheng/generate_material_list.py"), "exec"), namespace)
    wanted = {"insert_nozzle_flange_attachment_rows", "_first_empty_material_row", "_insert_component_row",
              "_remove_all_unused_components", "_component_source_candidates",
              "_is_empty_component_param_value", "_is_meaningful_component_source_value",
              "_finalize_material_name_specs", "clean_and_renumber"}
    path = ROOT / "modules/wenbenshengcheng/cunguige.py"
    tree = ast.parse(path.read_text(encoding="utf-8-sig"))
    exec(compile(ast.Module(body=[node for node in tree.body if isinstance(node, ast.FunctionDef)
                                and node.name in wanted], type_ignores=[]), str(path), "exec"), namespace)
    return namespace


class NozzleFlangeMaterialTests(unittest.TestCase):
    def test_each_title_group_uses_its_own_material_and_spec(self):
        rows = attachment(groups={
            PAIRED_FLANGE: {"材料类型": "钢锻件", "材料牌号": "16Mn", "锻件级别": "Ⅱ", "供货状态": "正火",
                            "配对法兰是否添加覆层": "是", "配对法兰覆层材料牌号": "S30403"},
            GASKET: {"垫片类型": "钢制管法兰用缠绕式垫片", "垫片型式": "D", "垫片材料": "1222"},
            BOLT: {"材料牌号": "35CrMo", "螺栓公称直径": "16", "螺栓长度": "85", "数量": "8"},
            NUT: {"材料牌号": "30CrMo", "数量": "16"},
        })
        calc = {"DictOutDatas": {"介质入口接管": {"Datas": [
            {"Name": "接管与管法兰或外部连接端壁厚（焊端规格）", "Value": "8"}]}}}
        bundle, = build_nozzle_flange_materials(rows, [nozzle()], calc)
        flange, gasket, bolt, nut = bundle["items"]
        self.assertEqual([item["material"] for item in bundle["items"]], ["16MnⅡ+S30403", "1222", "35CrMo", "30CrMo"])
        self.assertEqual([item["spec"] for item in bundle["items"]], ["WN 100-300 RF", "缠绕垫 D 100-300", "M16x85", "M16"])
        self.assertEqual([item["quantity"] for item in bundle["items"]], [1, 1, 8, 16])
        self.assertEqual(flange["supply"], "正火")
        self.assertTrue(all(item["name"].startswith("介质入口") for item in bundle["items"]))

    def test_multiple_nozzles_and_tabs_do_not_mix_same_named_parameters(self):
        rows = attachment("N1、N2，N1", groups={BOLT: {"材料牌号": "35CrMo"}})
        rows += attachment("N3", tab="配对2", tab_id="2", groups={BOLT: {"材料牌号": "25Cr2MoVA"}})
        bundles = build_nozzle_flange_materials(rows, [nozzle(), nozzle("N2", "气相出口"), nozzle("N3", "排污口")], {})
        self.assertEqual([bundle["code"] for bundle in bundles], ["N1", "N2", "N3"])
        self.assertEqual([bundle["items"][2]["material"] for bundle in bundles], ["35CrMo", "35CrMo", "25Cr2MoVA"])
        self.assertEqual(bundles[2]["items"][2]["name"], "排污口螺栓（接管法兰）")

    def test_incomplete_groups_still_create_four_names_without_inventing_dimensions(self):
        rows = attachment(groups={PAIRED_FLANGE: {"材料牌号": "16Mn", "材料类型": "钢锻件"}})
        bundle, = build_nozzle_flange_materials(rows, [nozzle()], {})
        self.assertEqual(len(bundle["items"]), 4)
        self.assertIsNone(bundle["items"][1]["spec"])
        self.assertIsNone(bundle["items"][2]["spec"])
        self.assertIsNone(bundle["items"][2]["quantity"])
        self.assertIsNone(bundle["items"][3]["quantity"])
        self.assertEqual(bundle["items"][0]["material"], "16Mn")

    def test_fullwidth_group_brackets_and_nps_match_existing_flange_format(self):
        rows = attachment(groups={"螺栓（接管法兰）": {"规格": "M16×85"}, GASKET: {"垫片型式": "D", "垫片类型": "缠绕式垫片"}})
        bundle, = build_nozzle_flange_materials(rows, [nozzle(size="4")], {})
        self.assertEqual(bundle["items"][0]["spec"], "WN 100-300 RF")
        self.assertEqual(bundle["items"][1]["spec"], "缠绕垫 D 100-300")
        self.assertEqual(bundle["items"][2]["spec"], "M16x85")
        self.assertEqual(bundle["items"][3]["spec"], "M16")

    def test_unrelated_attachment_and_unassigned_or_unknown_nozzles_add_no_rows(self):
        unrelated = attachment()
        unrelated[0]["附件类型"] = "防冲挡板"
        for rows in (unrelated, attachment(""), attachment("N99")):
            with self.subTest(rows=rows):
                self.assertEqual(build_nozzle_flange_materials(rows, [nozzle()], {}), [])

    def test_nozzle_gasket_uses_nominal_size_instead_of_equipment_gasket_dimensions(self):
        rows = attachment(groups={GASKET: {"垫片类型": "钢制管法兰用非金属平垫片", "垫片名义外径": "120",
                                         "垫片名义内径": "100", "垫片厚度": "3"}})
        bundle, = build_nozzle_flange_materials(rows, [nozzle()], {})
        self.assertEqual(bundle["items"][1]["spec"], "垫片 RF 100-300")

    def test_every_gasket_family_matches_requested_format_for_both_series(self):
        cases = [
            ("非金属平垫片", "", "垫片 RF 100-300"),
            ("聚四氟乙烯包覆垫片", "V", "四氟包覆垫 V 100-300"),
            ("金属包覆垫片", "I型", "金属包垫片 100-300"),
            ("缠绕式垫片", "D", "缠绕垫 D 100-300"),
            ("具有覆盖层的齿形组合垫", "B", "齿形垫 B 100-300"),
            ("金属环形垫", "椭圆形", "椭圆垫 100-300"),
            ("金属环形垫", "八角形", "八角垫 100-300"),
        ]
        for series in ("PN", "CLASS"):
            for family, style, expected in cases:
                with self.subTest(series=series, family=family, style=style):
                    rows = attachment(groups={GASKET: {"垫片类型": f"钢制管法兰用{family}({series}系列)", "垫片型式": style}})
                    bundle, = build_nozzle_flange_materials(rows, [nozzle()], {})
                    self.assertEqual(bundle["items"][1]["spec"], expected)

    def test_material_only_attachment_gets_bolt_and_nut_specs_from_standard_dimensions(self):
        rows = attachment(groups={BOLT: {"材料牌号": "SA-193 B8"}, NUT: {"材料牌号": "20"}})
        dimensions = {"N1": {"diameter": "16", "length": "95", "quantity": 8, "nut_quantity": 16}}
        bundle, = build_nozzle_flange_materials(rows, [nozzle()], {}, dimensions)
        self.assertEqual([item["spec"] for item in bundle["items"]][2:], ["M16x95", "M16"])
        self.assertEqual([item["quantity"] for item in bundle["items"]][2:], [8, 16])
        self.assertEqual(bundle["items"][2]["material"], "SA-193 B8")

    def test_explicit_stud_spec_keeps_its_name_and_overrides_standard_length(self):
        rows = attachment(groups={STUD: {"螺柱规格": "M16×120", "材料牌号": "35CrMo"}})
        dimensions = {"N1": {"diameter": "16", "length": "95"}}
        bundle, = build_nozzle_flange_materials(rows, [nozzle()], {}, dimensions)
        self.assertEqual(bundle["items"][2]["name"], "介质入口螺柱（接管法兰）")
        self.assertEqual(bundle["items"][2]["spec"], "M16x120")
        self.assertEqual(bundle["items"][3]["spec"], "M16")


class StandardFastenerTests(unittest.TestCase):
    class_table = "_20634_DNLess600法兰配用螺柱长度和质量CLASS"
    pn_table = "_20613_螺柱的长度和近似质量_带颈平焊_带颈对焊_螺纹以及承插焊法兰PN"

    def setUp(self):
        self.tables = {self.class_table: [
            {"DN": "100", "NPS": "4", "CLASS": "150", "PN": "20", "螺纹": "M16",
             "数量": "8", "LZR": "95", "LZJ": "105"},
            {"DN": "100", "NPS": "4", "CLASS": "300", "PN": "50", "螺纹": "M20",
             "数量": "8", "LZR": "125", "LZM": "130", "LZJ": "140"},
        ], self.pn_table: [{"DN": "50", "PN": "16", "螺纹": "M16", "数量": "4", "LZR": "85"}]}

    def test_class_and_nps_settings_select_correct_standard_row(self):
        port = nozzle(size="4", pressure="300")
        dims = resolve_nozzle_fastener_dimensions(port, {"公称尺寸类型": "NPS", "公称压力类型": "Class"}, self.tables)
        self.assertEqual((dims["diameter"], dims["length"]), ("20", "125"))

    def test_pn_is_matched_as_pn_not_class(self):
        port = nozzle(size="100", pressure="50")
        dims = resolve_nozzle_fastener_dimensions(port, {"公称压力类型": "PN"}, self.tables)
        self.assertEqual(dims["length"], "125")
        port = nozzle(size="50", pressure="16")
        port["法兰标准"] = "HG/T 20592-2009"
        self.assertEqual(resolve_nozzle_fastener_dimensions(port, {}, self.tables)["length"], "85")

    def test_nps_converts_to_dn_for_pn_tables_that_only_store_dn(self):
        port = nozzle(size="2", pressure="16")
        port["法兰标准"] = "HG/T 20592-2009"
        dims = resolve_nozzle_fastener_dimensions(port, {"公称尺寸类型": "NPS"}, self.tables)
        self.assertEqual((dims["diameter"], dims["length"]), ("16", "85"))

    def test_ring_joint_uses_its_own_length_not_raised_face_length(self):
        port = nozzle(pressure="150")
        port["密封面型式"] = "RJ"
        self.assertEqual(resolve_nozzle_fastener_dimensions(port, {}, self.tables)["length"], "105")

    def test_unmatched_size_pressure_and_face_do_not_reuse_first_standard_row(self):
        for change in ({"公称尺寸": "99"}, {"压力等级": "999"}, {"密封面型式": "TG"}, {"法兰标准": "其它标准"}):
            port = nozzle()
            port.update(change)
            with self.subTest(change=change):
                self.assertEqual(resolve_nozzle_fastener_dimensions(port, {}, self.tables), {})


class NozzleFlangeWorksheetTests(unittest.TestCase):
    def setUp(self):
        self.ns = worksheet_namespace()
        self.wb = openpyxl.load_workbook(ROOT / "templates/material_list/材料清单.xlsx")
        self.sheet = self.wb.active
        for row in self.ns["iter_material_rows"](self.sheet):
            for col in self.ns["MATERIAL_OUTPUT_COLUMNS"]:
                self.ns["_clear_material_cell"](row[col - 1])
        self.sheet.cell(7, 2, "介质入口接管法兰")
        self.sheet.cell(7, 25, "WN 100-300 RF s=8mm")
        self.sheet.cell(7, 15, 3.25)
        self.sheet.cell(8, 2, "气相出口接管法兰")
        self.sheet.cell(8, 25, "WN 80-300 RF s=8mm")

    def configure_db(self, rows, nozzles):
        cursor = MagicMock()
        cursor.fetchall.side_effect = [rows, nozzles]
        cursor.fetchone.return_value = {"公称尺寸类型": "DN", "公称压力类型": "Class"}
        conn = MagicMock()
        conn.cursor.return_value = cursor
        self.ns["get_connection"] = MagicMock(return_value=conn)
        cleanup_cursor = MagicMock()
        cleanup_cursor.fetchall.return_value = []
        cleanup_conn = MagicMock()
        cleanup_conn.cursor.return_value = cleanup_cursor
        cleanup_cursor.__enter__.return_value = cleanup_cursor
        self.ns["pymysql"] = SimpleNamespace(connect=MagicMock(return_value=cleanup_conn), cursors=SimpleNamespace(DictCursor=object))
        return conn, cursor

    def test_insert_clean_finalize_and_reopen_preserve_all_names_and_template_tail(self):
        conn, cursor = self.configure_db(attachment(groups={PAIRED_FLANGE: {"材料牌号": "16Mn"}}), [nozzle()])
        protected = self.ns["insert_nozzle_flange_attachment_rows"](self.sheet, "product1", {})
        self.ns["clean_and_renumber"](self.sheet, "product1")
        removed = self.ns["_remove_all_unused_components"](self.sheet, "product1", {}, protected)
        self.assertEqual(removed, 0)
        self.ns["_finalize_material_name_specs"](self.sheet)
        stream = BytesIO()
        self.wb.save(stream)
        stream.seek(0)
        reopened = openpyxl.load_workbook(stream).active
        self.assertEqual([reopened.cell(i, 2).value for i in range(7, 13)], [
            "介质入口接管法兰", "介质入口接管法兰配对法兰", "介质入口接管法兰垫片",
            "介质入口螺栓（接管法兰）", "介质入口螺母（接管法兰）", "气相出口接管法兰"])
        self.assertEqual(reopened.cell(8, 3).value, "WN 100-300 RF")
        self.assertEqual(reopened.cell(8, 15).value, 3.25)
        self.assertEqual(len(protected), 4)
        tail = self.ns["find_material_fixed_tail_start"](reopened)
        self.assertEqual(tuple(reopened.cell(tail + i, 2).value for i in range(4)), self.ns["MATERIAL_FIXED_TAIL_NAMES"])
        conn.close.assert_called_once()
        cursor.close.assert_called_once()
        self.assertIn("参数数值", cursor.execute.call_args_list[0].args[0])
        self.assertEqual(cursor.execute.call_args_list[0].args[1], ("product1", PAIRED_FLANGE))

    def test_overflow_inserts_each_bundle_after_its_own_flange(self):
        rows = attachment() + attachment("N2", "配对2", "2")
        self.configure_db(rows, [nozzle(), nozzle("N2", "气相出口", size="80")])
        for idx in range(9, self.ns["find_material_fixed_tail_start"](self.sheet)):
            self.sheet.cell(idx, 2, f"原有元件{idx}")
        old_tail = self.ns["find_material_fixed_tail_start"](self.sheet)
        self.ns["insert_nozzle_flange_attachment_rows"](self.sheet, "product1", {})
        names = [row[1].value for row in self.ns["iter_material_rows"](self.sheet)]
        self.assertEqual(names[0:2], ["介质入口接管法兰", "介质入口接管法兰配对法兰"])
        self.assertEqual(names[5:7], ["气相出口接管法兰", "气相出口接管法兰配对法兰"])
        self.assertEqual(self.ns["find_material_fixed_tail_start"](self.sheet), old_tail + 8)
        self.assertEqual(names[10], "原有元件9")

    def test_missing_attachment_does_not_change_material_rows(self):
        conn, cursor = self.configure_db([], [])
        before = list(self.sheet.values)
        self.assertEqual(self.ns["insert_nozzle_flange_attachment_rows"](self.sheet, "product1", {}), set())
        self.assertEqual(list(self.sheet.values), before)
        self.assertEqual(cursor.execute.call_count, 1)
        conn.close.assert_called_once()

    def test_name_only_exception_does_not_protect_unrelated_empty_components(self):
        self.configure_db(attachment(), [nozzle()])
        protected = self.ns["insert_nozzle_flange_attachment_rows"](self.sheet, "product1", {})
        self.sheet.cell(13, 2, "无效空垫片")
        self.assertEqual(self.ns["_remove_all_unused_components"](self.sheet, "product1", {}, protected), 1)
        self.assertTrue(protected.issubset({row[1].value for row in self.ns["iter_material_rows"](self.sheet)}))

    def test_standard_specs_reach_final_excel_spec_column_after_save_and_reopen(self):
        self.configure_db(attachment(groups={BOLT: {"材料牌号": "SA-193 B8"}}), [nozzle(pressure="150")])
        self.ns["load_nozzle_fastener_dimensions"].return_value = {
            "N1": {"diameter": "16", "length": "95", "quantity": 8, "nut_quantity": 16}}
        self.ns["insert_nozzle_flange_attachment_rows"](self.sheet, "product1", {})
        self.ns["_finalize_material_name_specs"](self.sheet)
        stream = BytesIO()
        self.wb.save(stream)
        stream.seek(0)
        saved = openpyxl.load_workbook(stream).active
        self.assertEqual(saved.cell(10, 3).value, "M16x95")
        self.assertEqual(saved.cell(11, 3).value, "M16")
        self.ns["load_nozzle_fastener_dimensions"].assert_called_once_with(
            [nozzle(pressure="150")], {"公称尺寸类型": "DN", "公称压力类型": "Class"})


class NozzleReinforcementRingTests(unittest.TestCase):
    def setUp(self):
        self.ns = worksheet_namespace()

    def sheet_with_names(self, names):
        wb = openpyxl.load_workbook(ROOT / "templates/material_list/材料清单.xlsx")
        sheet = wb.active
        for row in self.ns["iter_material_rows"](sheet):
            for col in self.ns["MATERIAL_OUTPUT_COLUMNS"]:
                self.ns["_clear_material_cell"](row[col - 1])
        for idx, name in enumerate(names, start=7):
            self.ns["ensure_material_row_capacity"](sheet, idx)
            sheet.cell(idx, 2, name)
            sheet.cell(idx, 5, 1)
            sheet.cell(idx, 15, 12.5)
            sheet.cell(idx, 25, "Ø180/Ø114；δ=8")
        return wb, sheet

    @staticmethod
    def calculation(**flags):
        return {"DictOutDatas": {name + "接管": {"Datas": [
            {"Name": "是否使用补强圈", "Value": value},
            {"Name": "接管序号", "Value": code},
            {"Name": "补强圈面积", "Value": "500"},
        ]} for code, (name, value) in flags.items()}}

    def names(self, sheet):
        return [row[1].value for row in self.ns["iter_material_rows"](sheet) if row[1].value]

    def test_every_value_other_than_yes_removes_ring_despite_old_positive_dimensions_and_mass(self):
        for value in ("否", "", None, "程序推荐", 0, False, "yes", "是（推荐）"):
            with self.subTest(value=value):
                _, sheet = self.sheet_with_names(["介质入口补强圈", "介质入口接管"])
                data = self.calculation(N1=("介质入口", value))
                self.assertEqual(self.ns["remove_unused_nozzle_reinforcement_rings"](sheet, data), 1)
                self.assertEqual(self.names(sheet), ["介质入口接管"])

    def test_each_ring_follows_its_own_nozzle_flag_and_keeps_unrelated_components(self):
        _, sheet = self.sheet_with_names([
            "介质入口补强圈", "气相出口接管补强圈", "补强圈（N3）", "N1补强圈", "N2补强圈",
            "介质入口接管", "气相出口接管法兰", "加强圈",
        ])
        data = self.calculation(N1=("介质入口", "是"), N2=("气相出口", "否"))
        self.assertEqual(self.ns["remove_unused_nozzle_reinforcement_rings"](sheet, data), 3)
        self.assertEqual(self.names(sheet), ["介质入口补强圈", "N1补强圈", "介质入口接管", "气相出口接管法兰", "加强圈"])

    def test_missing_switch_or_missing_nozzle_output_removes_ring(self):
        for data in ({}, {"DictOutDatas": {"介质入口接管": {"Datas": [
            {"Name": "补强圈面积", "Value": "500"}]}}}):
            with self.subTest(data=data):
                _, sheet = self.sheet_with_names(["介质入口补强圈"])
                self.assertEqual(self.ns["remove_unused_nozzle_reinforcement_rings"](sheet, data), 1)

    def test_yes_is_trimmed_and_legacy_name_with_spec_still_matches_own_nozzle(self):
        _, sheet = self.sheet_with_names(["介质入口补强圈Ø180/Ø114；δ=8", "气相出口补强圈 Ø180/Ø114；δ=8"])
        data = self.calculation(N1=("介质入口", " 是 "), N2=("气相出口", "否"))
        self.assertEqual(self.ns["remove_unused_nozzle_reinforcement_rings"](sheet, data), 1)
        self.assertEqual(self.names(sheet), ["介质入口补强圈Ø180/Ø114；δ=8"])

    def test_summary_ring_is_removed_only_when_no_nozzle_uses_one(self):
        for flag, removed in (("否", 1), ("是", 0)):
            with self.subTest(flag=flag):
                _, sheet = self.sheet_with_names(["补强圈"])
                self.assertEqual(self.ns["remove_unused_nozzle_reinforcement_rings"](
                    sheet, self.calculation(N1=("介质入口", flag))), removed)

    def test_final_cleanup_enforces_ring_switch_before_preserve_list_and_save(self):
        wb, sheet = self.sheet_with_names(["介质入口补强圈", "气相出口补强圈"])
        cursor = MagicMock()
        cursor.fetchall.return_value = []
        conn = MagicMock()
        conn.cursor.return_value = cursor
        self.ns["pymysql"] = SimpleNamespace(connect=MagicMock(return_value=conn), cursors=SimpleNamespace(DictCursor=object))
        data = self.calculation(N1=("介质入口", "否"), N2=("气相出口", "是"))
        removed = self.ns["_remove_all_unused_components"](
            sheet, "product1", data, preserved_names={"介质入口补强圈"})
        self.assertEqual(removed, 1)
        self.ns["_finalize_material_name_specs"](sheet)
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        saved = openpyxl.load_workbook(stream).active
        self.assertEqual(self.names(saved), ["气相出口补强圈"])
        tail = self.ns["find_material_fixed_tail_start"](saved)
        self.assertEqual(tuple(saved.cell(tail + i, 2).value for i in range(4)), self.ns["MATERIAL_FIXED_TAIL_NAMES"])


if __name__ == "__main__":
    unittest.main()
