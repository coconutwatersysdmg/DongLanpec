import os
import re
import json
import threading
from typing import List, Dict, Any, Optional, Tuple
import tkinter.font as tkfont

import pdfplumber
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# =========================
# PDF 表格解析核心（按整张表输出）
# =========================
class PDFTableParser:
    def __init__(self, line_tol: float = 3.0):
        self.line_tol = line_tol

    def parse_pdf(self, pdf_path: str) -> Dict[str, Any]:
        result = {
            "source_pdf": os.path.abspath(pdf_path),
            "pages": []
        }

        with pdfplumber.open(pdf_path) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                tables = self._parse_page_tables(page, page_no)
                result["pages"].append({
                    "page": page_no,
                    "width": page.width,
                    "height": page.height,
                    "tables": tables
                })

        return result

    def _parse_page_tables(self, page, page_no: int) -> List[Dict[str, Any]]:
        table_settings = {
            "vertical_strategy": "lines",
            "horizontal_strategy": "lines",
            "intersection_tolerance": self.line_tol,
            "snap_tolerance": self.line_tol,
            "join_tolerance": self.line_tol,
            "edge_min_length": 8,
            "min_words_vertical": 1,
            "min_words_horizontal": 1,
        }

        found_tables = page.find_tables(table_settings=table_settings)
        results = []

        for idx, table in enumerate(found_tables, start=1):
            raw = table.extract()
            if not raw:
                continue

            data = self._normalize_2d(raw)
            if not self._has_meaningful_content(data):
                continue

            bbox = list(table.bbox) if table.bbox else None
            title = self._guess_table_title(page, data, bbox)
            orientation = self._guess_orientation(data)
            table_type = self._guess_table_type(title, data)
            headers, body = self._split_header_body(data)
            headers = self._compact_rows_for_output(headers, table_type)
            body = self._compact_rows_for_output(body, table_type)
            checks = self._extract_checks(data)
            notes = self._extract_notes(data)
            subtables = self._extract_subtables_from_data(body, page_no, idx)

            results.append({
                "table_id": f"p{page_no}_t{idx}",
                "title": title,
                "type": table_type,
                "bbox": bbox,
                "orientation": orientation,
                "rows": len(data),
                "cols": max((len(r) for r in data), default=0),
                "headers": headers,
                "data": body,
                "subtables": subtables,
                "checks": checks,
                "notes": notes,
            })

        return self._merge_nested_tables(results)

    def _normalize_2d(self, raw: List[List[Any]]) -> List[List[str]]:
        width = max((len(r) for r in raw), default=0)
        normalized = []
        for row in raw:
            new_row = [self._clean_text(c) for c in row]
            if len(new_row) < width:
                new_row.extend([""] * (width - len(new_row)))
            normalized.append(new_row)
        return normalized

    def _remove_empty_edges(self, data: List[List[str]]) -> List[List[str]]:
        while data and all(not c.strip() for c in data[0]):
            data.pop(0)
        while data and all(not c.strip() for c in data[-1]):
            data.pop()
        if not data:
            return []

        cols = len(data[0])
        left = 0
        right = cols - 1

        while left < cols and all((row[left].strip() == "") for row in data):
            left += 1
        while right >= 0 and all((row[right].strip() == "") for row in data):
            right -= 1

        if left > right:
            return []
        return [row[left:right + 1] for row in data]

    def _trim_row_trailing_empty(self, data: List[List[str]]) -> List[List[str]]:
        trimmed = []
        for row in data:
            new_row = row[:]
            while new_row and new_row[-1] == "":
                new_row.pop()
            trimmed.append(new_row)
        return trimmed

    def _compress_sparse_rows(self, data: List[List[str]]) -> List[List[str]]:
        compressed = []
        for row in data:
            non_empty_count = sum(1 for c in row if c.strip())
            if non_empty_count == 0:
                continue

            if non_empty_count <= max(2, len(row) // 3):
                new_row = [c for c in row if c.strip()]
                compressed.append(new_row)
            else:
                compressed.append(row)
        return compressed

    def _has_meaningful_content(self, data: List[List[str]]) -> bool:
        non_empty = sum(1 for row in data for c in row if c.strip())
        return non_empty >= 2

    def _guess_table_title(self, page, data: List[List[str]], bbox: Optional[List[float]]) -> str:
        for row in data[:3]:
            joined = " ".join([c for c in row if c]).strip()
            if len(joined) >= 4:
                return joined[:120]

        if bbox:
            x0, top, x1, _ = bbox
            search_top = max(0, top - 45)
            words = page.extract_words()
            title_words = []
            for w in words:
                wx0 = w.get("x0", 0)
                wx1 = w.get("x1", 0)
                wtop = w.get("top", 0)
                if wx0 >= x0 - 5 and wx1 <= x1 + 5 and search_top <= wtop <= top:
                    title_words.append(w.get("text", ""))
            title = " ".join(title_words).strip()
            if title:
                return title[:120]

        return "未命名表格"

    def _guess_orientation(self, data: List[List[str]]) -> str:
        rows = len(data)
        cols = max((len(r) for r in data), default=0)
        if rows == 0 or cols == 0:
            return "unknown"
        if rows > cols * 1.4:
            return "vertical"
        if cols > rows * 1.4:
            return "horizontal"
        return "mixed"

    def _guess_table_type(self, title: str, data: List[List[str]]) -> str:
        text = (title + " " + " ".join(" ".join(r) for r in data[:5])).lower()
        if "工艺" in text or "process" in text:
            return "process_table"
        if "结构" in text or "structure" in text:
            return "structure_table"
        if "设计" in text or "design" in text:
            return "design_table"
        if "基本" in text or "设备" in text or "位号" in text:
            return "info_table"
        return "generic_table"

    def _split_header_body(self, data: List[List[str]]) -> Tuple[List[List[str]], List[List[str]]]:
        if not data:
            return [], []

        header_rows = 1
        if len(data) >= 2:
            first_score = self._header_score(data[0])
            second_score = self._header_score(data[1])
            if second_score >= 2 and first_score >= 2:
                header_rows = 2

        headers = data[:header_rows]
        body = data[header_rows:] if len(data) > header_rows else []
        return headers, body

    def _compact_rows_for_output(self, rows: List[List[str]], table_type: str) -> List[List[str]]:
        return rows

    def _compact_single_row(self, row: List[str], table_type: str) -> List[str]:
        return row

    def _header_score(self, row: List[str]) -> int:
        score = 0
        header_keywords = [
            "名称", "name", "单位", "unit", "项目", "item", "参数", "数值",
            "壳程", "管程", "shell", "tube", "进口", "出口", "类型", "规格"
        ]
        for c in row:
            t = c.strip().lower()
            if any(k in t for k in header_keywords):
                score += 1
        return score

    def _extract_checks(self, data: List[List[str]]) -> List[Dict[str, Any]]:
        checks = []
        positive_marks = ["☑", "☒", "√", "✓", "■"]
        negative_marks = ["☐", "□"]

        for row in data:
            for cell in row:
                if not cell:
                    continue
                value = None
                if any(m in cell for m in positive_marks):
                    value = True
                elif any(m in cell for m in negative_marks):
                    value = False
                elif re.search(r"(^|\s)(是|有)(\s|$)", cell):
                    value = True
                elif re.search(r"(^|\s)(否|无)(\s|$)", cell):
                    value = False

                if value is not None:
                    label = re.sub(r"[☑☒√✓■☐□]", "", cell).strip()
                    checks.append({
                        "label": label or cell,
                        "value": value
                    })
        return checks

    def _extract_notes(self, data: List[List[str]]) -> List[str]:
        notes = []
        for row in data:
            joined = " ".join([c for c in row if c]).strip()
            if not joined:
                continue
            if any(k in joined for k in ["注", "备注", "note", "说明"]):
                notes.append(joined)
        return notes

    def _extract_subtables_from_data(self, data: List[List[str]], page_no: int, table_idx: int) -> List[Dict[str, Any]]:
        subtables = []
        current_title = None
        current_rows = []
        sub_idx = 0

        for row in data:
            non_empty = [c for c in row if c.strip()]
            joined = " ".join(non_empty).strip()

            if self._looks_like_subtable_title(joined):
                if current_rows:
                    sub_idx += 1
                    subtables.append({
                        "subtable_id": f"p{page_no}_t{table_idx}_s{sub_idx}",
                        "title": current_title or f"子表{sub_idx}",
                        "headers": [current_rows[0]] if current_rows else [],
                        "data": current_rows[1:] if len(current_rows) > 1 else [],
                    })
                    current_rows = []
                current_title = joined
            elif current_title:
                current_rows.append(row)

        if current_title and current_rows:
            sub_idx += 1
            subtables.append({
                "subtable_id": f"p{page_no}_t{table_idx}_s{sub_idx}",
                "title": current_title,
                "headers": [current_rows[0]] if current_rows else [],
                "data": current_rows[1:] if len(current_rows) > 1 else [],
            })

        return subtables

    def _looks_like_subtable_title(self, text: str) -> bool:
        if not text or len(text) > 40:
            return False
        keywords = ["工艺操作数据", "结构参数", "设计数据", "材料", "要求", "参数", "数据表"]
        return any(k in text for k in keywords)

    def _merge_nested_tables(self, tables: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        if not tables:
            return tables

        merged = []
        used = set()
        for i, parent in enumerate(tables):
            if i in used:
                continue
            pb = parent.get("bbox")
            if not pb:
                merged.append(parent)
                continue

            parent_subtables = list(parent.get("subtables", []))
            for j, child in enumerate(tables):
                if i == j or j in used:
                    continue
                cb = child.get("bbox")
                if not cb:
                    continue
                if self._is_inside(cb, pb):
                    parent_subtables.append({
                        "subtable_id": child.get("table_id", f"sub_{j+1}"),
                        "title": child.get("title", "子表"),
                        "type": child.get("type", "generic_table"),
                        "headers": child.get("headers", []),
                        "data": child.get("data", []),
                        "checks": child.get("checks", []),
                        "notes": child.get("notes", []),
                    })
                    used.add(j)

            parent["subtables"] = parent_subtables
            merged.append(parent)

        final_tables = []
        for idx, t in enumerate(merged):
            if idx not in used:
                final_tables.append(t)
        return final_tables

    def _is_inside(self, inner: List[float], outer: List[float], tol: float = 3.0) -> bool:
        return (
            inner[0] >= outer[0] - tol and
            inner[1] >= outer[1] - tol and
            inner[2] <= outer[2] + tol and
            inner[3] <= outer[3] + tol and
            (inner != outer)
        )

    @staticmethod
    def _clean_text(text: Any) -> str:
        if text is None:
            return ""
        return " ".join(str(text).replace("\n", " ").split())


# =========================
# 聊天气泡日志组件
# =========================
class ChatLog(tk.Frame):
    def __init__(self, master, bg="#ffffff", name_font=None, text_font=None):
        super().__init__(master, bg=bg)

        self.name_font = name_font or tkfont.Font(family="微软雅黑", size=9)
        self.text_font = text_font or tkfont.Font(family="微软雅黑", size=10)
        self.base_wraplength = 640
        self.current_wraplength = 640

        self.canvas = tk.Canvas(
            self,
            bg=bg,
            highlightthickness=0,
            bd=0
        )
        self.scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.inner = tk.Frame(self.canvas, bg=bg)

        self.inner.bind(
            "<Configure>",
            lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )

        self.window_id = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.canvas.bind(
            "<Configure>",
            lambda e: self.canvas.itemconfig(self.window_id, width=e.width)
        )
        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")

    def set_scale(self, scale: float):
        self.current_wraplength = int(self.base_wraplength * scale)

        for child in self.inner.winfo_children():
            for sub in child.winfo_children():
                try:
                    if isinstance(sub, tk.Label):
                        role = getattr(sub, "_chat_role", "")
                        if role == "name":
                            sub.configure(font=self.name_font)
                        elif role == "bubble":
                            sub.configure(font=self.text_font, wraplength=self.current_wraplength)
                except Exception:
                    pass

    def add_message(self, text: str, sender: str = "system"):
        wrap_len = self.current_wraplength

        row = tk.Frame(self.inner, bg="#ffffff")
        row.pack(fill="x", padx=14, pady=8, anchor="w")

        if sender == "user":
            bubble_bg = "#e8f1ff"
            bubble_fg = "#1d4ed8"
            anchor = "e"
            side = "right"
            name = "你"
        elif sender == "success":
            bubble_bg = "#ecfdf5"
            bubble_fg = "#047857"
            anchor = "w"
            side = "left"
            name = "系统"
        elif sender == "error":
            bubble_bg = "#fef2f2"
            bubble_fg = "#b91c1c"
            anchor = "w"
            side = "left"
            name = "系统"
        else:
            bubble_bg = "#f3f4f6"
            bubble_fg = "#374151"
            anchor = "w"
            side = "left"
            name = "系统"

        name_label = tk.Label(
            row,
            text=name,
            bg="#ffffff",
            fg="#9ca3af",
            font=self.name_font
        )
        name_label._chat_role = "name"
        name_label.pack(anchor=anchor, padx=6, pady=(0, 4))

        bubble = tk.Label(
            row,
            text=text,
            bg=bubble_bg,
            fg=bubble_fg,
            justify="left",
            wraplength=wrap_len,
            padx=14,
            pady=10,
            font=self.text_font,
            relief="flat"
        )
        bubble._chat_role = "bubble"
        bubble.pack(side=side, anchor=anchor, padx=6)

        self.update_idletasks()
        self.canvas.yview_moveto(1.0)

    def clear(self):
        for child in self.inner.winfo_children():
            child.destroy()
        self.update_idletasks()
        self.canvas.yview_moveto(0.0)


# =========================
# GUI
# =========================
class App:
    def __init__(self, root):
        self.root = root
        self.root.title("PDF -> JSON / Excel")
        self.root.geometry("1280x860")
        self.root.minsize(1120, 720)
        self.root.configure(bg="#ffffff")

        self.parser = PDFTableParser()
        self.input_mode = tk.StringVar(value="file")
        self.input_path = tk.StringVar()
        self.output_dir = tk.StringVar()
        self.status_text = tk.StringVar(value="就绪")
        self.export_excel_after_parse = tk.BooleanVar(value=True)

        self._last_font_scale = None

        self._init_fonts()
        self._configure_style()
        self._build_ui()

        self.root.bind("<Configure>", self._on_window_resize)

    def _init_fonts(self):
        self.font_title = tkfont.Font(family="微软雅黑", size=20, weight="bold")
        self.font_subtitle = tkfont.Font(family="微软雅黑", size=10)
        self.font_card_title = tkfont.Font(family="微软雅黑", size=11, weight="bold")
        self.font_card_text = tkfont.Font(family="微软雅黑", size=10)
        self.font_button = tkfont.Font(family="微软雅黑", size=10)
        self.font_radio = tkfont.Font(family="微软雅黑", size=10)
        self.font_status = tkfont.Font(family="微软雅黑", size=10, weight="bold")
        self.font_chat_name = tkfont.Font(family="微软雅黑", size=9)
        self.font_chat_text = tkfont.Font(family="微软雅黑", size=10)
        self.font_hint_title = tkfont.Font(family="微软雅黑", size=10, weight="bold")
        self.font_hint_text = tkfont.Font(family="微软雅黑", size=10)
        self.font_right_title = tkfont.Font(family="微软雅黑", size=12, weight="bold")

    def _on_window_resize(self, event=None):
        if event is not None and event.widget != self.root:
            return

        w = self.root.winfo_width()
        scale = w / 1280.0
        scale = max(0.9, min(scale, 1.8))

        rounded_scale = round(scale, 2)
        if self._last_font_scale == rounded_scale:
            return
        self._last_font_scale = rounded_scale

        self.font_title.configure(size=max(16, int(20 * scale)))
        self.font_subtitle.configure(size=max(9, int(10 * scale)))
        self.font_card_title.configure(size=max(10, int(11 * scale)))
        self.font_card_text.configure(size=max(9, int(10 * scale)))
        self.font_button.configure(size=max(9, int(10 * scale)))
        self.font_radio.configure(size=max(9, int(10 * scale)))
        self.font_status.configure(size=max(9, int(10 * scale)))
        self.font_chat_name.configure(size=max(8, int(9 * scale)))
        self.font_chat_text.configure(size=max(9, int(10 * scale)))
        self.font_hint_title.configure(size=max(9, int(10 * scale)))
        self.font_hint_text.configure(size=max(9, int(10 * scale)))
        self.font_right_title.configure(size=max(10, int(12 * scale)))

        if hasattr(self, "chat"):
            self.chat.set_scale(scale)

    def _configure_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass

        self.bg_main = "#ffffff"
        self.bg_panel = "#ffffff"
        self.bg_panel_2 = "#f8fafc"
        self.border = "#e5e7eb"
        self.text_main = "#111827"
        self.text_sub = "#6b7280"
        self.accent = "#2563eb"

        style.configure("App.TFrame", background=self.bg_main)
        style.configure("Panel.TFrame", background=self.bg_panel)
        style.configure("Inner.TFrame", background=self.bg_panel_2)

        style.configure(
            "Title.TLabel",
            background=self.bg_main,
            foreground=self.text_main,
            font=self.font_title
        )
        style.configure(
            "SubTitle.TLabel",
            background=self.bg_main,
            foreground=self.text_sub,
            font=self.font_subtitle
        )
        style.configure(
            "CardTitle.TLabel",
            background=self.bg_panel,
            foreground=self.text_main,
            font=self.font_card_title
        )
        style.configure(
            "CardText.TLabel",
            background=self.bg_panel,
            foreground=self.text_sub,
            font=self.font_card_text
        )

        style.configure(
            "Modern.TButton",
            font=self.font_button,
            padding=(14, 8),
            borderwidth=0
        )
        style.map(
            "Modern.TButton",
            background=[("active", "#1d4ed8"), ("!disabled", self.accent)],
            foreground=[("!disabled", "white")]
        )

        style.configure(
            "Ghost.TButton",
            font=self.font_button,
            padding=(12, 8),
            borderwidth=0
        )
        style.map(
            "Ghost.TButton",
            background=[("active", "#dbeafe"), ("!disabled", "#eff6ff")],
            foreground=[("!disabled", "#1d4ed8")]
        )

        style.configure(
            "Light.TLabelframe",
            background=self.bg_panel,
            foreground=self.text_main,
            bordercolor=self.border,
            relief="solid"
        )
        style.configure(
            "Light.TLabelframe.Label",
            background=self.bg_panel,
            foreground=self.text_main,
            font=self.font_card_title
        )

        style.configure(
            "Light.TRadiobutton",
            background=self.bg_panel,
            foreground=self.text_main,
            font=self.font_radio
        )
        style.map(
            "Light.TRadiobutton",
            background=[("active", self.bg_panel)],
            foreground=[("active", self.text_main)]
        )

        style.configure(
            "Light.TCheckbutton",
            background=self.bg_panel,
            foreground=self.text_main,
            font=self.font_radio
        )

        style.configure(
            "Light.TEntry",
            fieldbackground="#ffffff",
            background="#ffffff",
            foreground=self.text_main,
            bordercolor=self.border,
            lightcolor=self.border,
            darkcolor=self.border,
            padding=8
        )

        style.configure(
            "Blue.Horizontal.TProgressbar",
            troughcolor="#eef2ff",
            background=self.accent,
            bordercolor="#eef2ff",
            lightcolor=self.accent,
            darkcolor=self.accent
        )

    def _build_ui(self):
        main = ttk.Frame(self.root, style="App.TFrame", padding=18)
        main.pack(fill="both", expand=True)

        header = ttk.Frame(main, style="App.TFrame")
        header.pack(fill="x", pady=(0, 14))

        left_header = ttk.Frame(header, style="App.TFrame")
        left_header.pack(side="left", fill="y")

        ttk.Label(left_header, text="PDF 转 JSON / Excel", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            left_header,
            text="专注于 PDF 表格，并支持 JSON / Excel 导出",
            style="SubTitle.TLabel"
        ).pack(anchor="w", pady=(4, 0))

        status_wrap = tk.Frame(header, bg=self.bg_main)
        status_wrap.pack(side="right", anchor="ne", pady=4)

        self.status_badge = tk.Label(
            status_wrap,
            textvariable=self.status_text,
            bg="#eff6ff",
            fg="#1d4ed8",
            font=self.font_status,
            padx=14,
            pady=6
        )
        self.status_badge.pack()

        content = ttk.Frame(main, style="App.TFrame")
        content.pack(fill="both", expand=True)

        content.columnconfigure(0, weight=0)
        content.columnconfigure(1, weight=1)
        content.rowconfigure(0, weight=1)

        left_panel = tk.Frame(
            content,
            bg=self.bg_panel,
            highlightbackground=self.border,
            highlightthickness=1
        )
        left_panel.grid(row=0, column=0, sticky="nsw", padx=(0, 14))
        left_panel.configure(width=410)
        left_panel.grid_propagate(False)

        left_canvas = tk.Canvas(
            left_panel,
            bg=self.bg_panel,
            highlightthickness=0,
            bd=0
        )
        left_scrollbar = ttk.Scrollbar(left_panel, orient="vertical", command=left_canvas.yview)
        left_canvas.configure(yscrollcommand=left_scrollbar.set)

        left_scrollbar.pack(side="right", fill="y")
        left_canvas.pack(side="left", fill="both", expand=True)

        left_inner = ttk.Frame(left_canvas, style="Panel.TFrame")
        left_canvas_window = left_canvas.create_window((0, 0), window=left_inner, anchor="nw")

        def _on_left_inner_configure(event):
            left_canvas.configure(scrollregion=left_canvas.bbox("all"))

        def _on_left_canvas_configure(event):
            left_canvas.itemconfig(left_canvas_window, width=event.width)

        left_inner.bind("<Configure>", _on_left_inner_configure)
        left_canvas.bind("<Configure>", _on_left_canvas_configure)

        # 鼠标滚轮
        def _on_mousewheel(event):
            left_canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        left_canvas.bind_all("<MouseWheel>", _on_mousewheel)

        # 给 left_inner 保留原来的内边距容器
        inner_wrap = ttk.Frame(left_inner, style="Panel.TFrame")
        inner_wrap.pack(fill="both", expand=True, padx=18, pady=18)

        left_inner = inner_wrap
        ttk.Label(left_inner, text="控制面板", style="CardTitle.TLabel").pack(anchor="w")
        ttk.Label(
            left_inner,
            text="设置输入源、输出位置与执行方式。",
            style="CardText.TLabel"
        ).pack(anchor="w", pady=(4, 16))

        mode_card = ttk.LabelFrame(left_inner, text="输入方式", style="Light.TLabelframe")
        mode_card.pack(fill="x", pady=(0, 12), ipady=6)

        ttk.Radiobutton(
            mode_card, text="单个 PDF", variable=self.input_mode, value="file", style="Light.TRadiobutton"
        ).pack(anchor="w", padx=12, pady=(8, 4))

        ttk.Radiobutton(
            mode_card, text="文件夹批量处理", variable=self.input_mode, value="folder", style="Light.TRadiobutton"
        ).pack(anchor="w", padx=12, pady=(0, 8))

        input_card = ttk.LabelFrame(left_inner, text="PDF 或文件夹路径", style="Light.TLabelframe")
        input_card.pack(fill="x", pady=(0, 12), ipady=8)

        self.input_entry = ttk.Entry(input_card, textvariable=self.input_path, style="Light.TEntry")
        self.input_entry.pack(fill="x", padx=12, pady=(10, 8))

        ttk.Button(
            input_card, text="选择输入", style="Ghost.TButton", command=self.choose_input
        ).pack(anchor="w", padx=12, pady=(0, 8))

        out_card = ttk.LabelFrame(left_inner, text="输出目录", style="Light.TLabelframe")
        out_card.pack(fill="x", pady=(0, 12), ipady=8)

        self.output_entry = ttk.Entry(out_card, textvariable=self.output_dir, style="Light.TEntry")
        self.output_entry.pack(fill="x", padx=12, pady=(10, 8))

        ttk.Button(
            out_card, text="选择输出目录", style="Ghost.TButton", command=self.choose_output
        ).pack(anchor="w", padx=12, pady=(0, 8))

        hint_card = tk.Frame(
            left_inner,
            bg=self.bg_panel_2,
            highlightbackground=self.border,
            highlightthickness=1
        )
        hint_card.pack(fill="x", pady=(0, 14))

        tk.Label(
            hint_card,
            text="当前输出结构",
            bg=self.bg_panel_2,
            fg=self.text_main,
            font=self.font_hint_title
        ).pack(anchor="w", padx=12, pady=(12, 6))

        tk.Label(
            hint_card,
            text=(
                "• 按 页 → 表 的层级输出\n"
                "• 保留 title / headers / data / subtables / checks / notes\n"
                "• 保留空单元格，不压缩原始表格结构\n"
                "• 支持直接从 JSON 导出 Excel\n"
                "• 支持 PDF 解析后同步导出 Excel"
            ),
            justify="left",
            bg=self.bg_panel_2,
            fg=self.text_sub,
            font=self.font_hint_text,
            padx=12
        ).pack(anchor="w", fill="x", pady=(0, 12))

        action_wrap = ttk.Frame(left_inner, style="Panel.TFrame")
        action_wrap.pack(fill="x", pady=(0, 12))

        action_wrap.columnconfigure(0, weight=1)
        action_wrap.columnconfigure(1, weight=1)

        self.start_btn = ttk.Button(
            action_wrap, text="开始解析", style="Modern.TButton", command=self.start_parse
        )
        self.start_btn.grid(row=0, column=0, sticky="ew", padx=(0, 6), pady=(0, 8))

        self.clear_btn = ttk.Button(
            action_wrap, text="清空日志", style="Ghost.TButton", command=self.clear_log
        )
        self.clear_btn.grid(row=0, column=1, sticky="ew", padx=(6, 0), pady=(0, 8))

        self.json_to_excel_btn = ttk.Button(
            action_wrap, text="JSON转Excel", style="Ghost.TButton", command=self.choose_json_and_export_excel
        )
        self.json_to_excel_btn.grid(row=1, column=0, columnspan=2, sticky="ew")

        ttk.Checkbutton(
            left_inner,
            text="解析完成后同时导出 Excel",
            variable=self.export_excel_after_parse,
            style="Light.TCheckbutton"
        ).pack(anchor="w", pady=(4, 10))

        self.progress = ttk.Progressbar(
            left_inner,
            mode="indeterminate",
            style="Blue.Horizontal.TProgressbar"
        )
        self.progress.pack(fill="x", pady=(6, 12))

        right_panel = tk.Frame(
            content,
            bg=self.bg_panel,
            highlightbackground=self.border,
            highlightthickness=1
        )
        right_panel.grid(row=0, column=1, sticky="nsew")

        right_header = tk.Frame(right_panel, bg=self.bg_panel)
        right_header.pack(fill="x", padx=18, pady=(16, 10))

        tk.Label(
            right_header,
            text="运行日志",
            bg=self.bg_panel,
            fg=self.text_main,
            font=self.font_right_title
        ).pack(anchor="w")

        tk.Label(
            right_header,
            text="解析过程和导出结果会显示在这里。",
            bg=self.bg_panel,
            fg=self.text_sub,
            font=self.font_card_text
        ).pack(anchor="w", pady=(4, 0))

        self.chat = ChatLog(
            right_panel,
            bg="#ffffff",
            name_font=self.font_chat_name,
            text_font=self.font_chat_text
        )
        self.chat.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        self.chat.add_message("欢迎使用 PDF 解析工具。请先选择输入和输出路径。", "system")

    def choose_input(self):
        if self.input_mode.get() == "file":
            path = filedialog.askopenfilename(filetypes=[("PDF Files", "*.pdf")])
        else:
            path = filedialog.askdirectory()
        if path:
            self.input_path.set(path)
            self.chat.add_message(f"已选择输入：{path}", "user")

    def choose_output(self):
        path = filedialog.askdirectory()
        if path:
            self.output_dir.set(path)
            self.chat.add_message(f"已选择输出目录：{path}", "user")

    def clear_log(self):
        self.chat.clear()
        self.chat.add_message("日志已清空。", "system")

    def choose_json_and_export_excel(self):
        json_path = filedialog.askopenfilename(filetypes=[("JSON Files", "*.json")])
        if not json_path:
            return

        out_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile=os.path.splitext(os.path.basename(json_path))[0] + ".xlsx"
        )
        if not out_path:
            return

        try:
            with open(json_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            self.export_json_to_excel(data, out_path)
            self.chat.add_message(f"JSON 已导出 Excel：{out_path}", "success")
            messagebox.showinfo("完成", f"已导出 Excel：\n{out_path}")
        except Exception as e:
            self.chat.add_message(f"JSON 转 Excel 失败：{e}", "error")
            messagebox.showerror("错误", str(e))

    def start_parse(self):
        in_path = self.input_path.get().strip()
        out_dir = self.output_dir.get().strip()

        if not in_path:
            messagebox.showerror("错误", "请选择 PDF 文件或文件夹。")
            return
        if not out_dir:
            messagebox.showerror("错误", "请选择输出目录。")
            return

        self._set_running(True)
        threading.Thread(target=self._run_parse, daemon=True).start()

    def _run_parse(self):
        try:
            self._set_status("解析中...", badge_bg="#fef3c7", badge_fg="#92400e")
            mode = self.input_mode.get()
            in_path = self.input_path.get().strip()
            out_dir = self.output_dir.get().strip()

            pdf_files = []
            if mode == "file":
                pdf_files = [in_path]
            else:
                for root_dir, _, files in os.walk(in_path):
                    for f in files:
                        if f.lower().endswith(".pdf"):
                            pdf_files.append(os.path.join(root_dir, f))

            if not pdf_files:
                self._append_chat("未找到 PDF 文件。", "error")
                self._set_status("完成", badge_bg="#eff6ff", badge_fg="#1d4ed8")
                self._set_running(False)
                return

            self._append_chat(f"本次共检测到 {len(pdf_files)} 个 PDF 文件，开始处理。", "system")

            for index, pdf in enumerate(pdf_files, start=1):
                self._append_chat(f"[{index}/{len(pdf_files)}] 开始处理：{pdf}", "system")
                data = self.parser.parse_pdf(pdf)

                base_name = os.path.splitext(os.path.basename(pdf))[0]

                json_out_path = os.path.join(out_dir, base_name + ".json")
                with open(json_out_path, "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)

                excel_msg = ""
                if self.export_excel_after_parse.get():
                    excel_out_path = os.path.join(out_dir, base_name + ".xlsx")
                    self.export_json_to_excel(data, excel_out_path)
                    excel_msg = f"\n同时导出 Excel：{excel_out_path}"

                pages_count = len(data.get("pages", []))
                table_count = sum(len(p.get("tables", [])) for p in data.get("pages", []))
                self._append_chat(
                    f"已输出 JSON：{json_out_path}{excel_msg}\n共解析 {pages_count} 页，识别到 {table_count} 张表。",
                    "success"
                )

            self._set_status("解析完成", badge_bg="#dcfce7", badge_fg="#166534")
            self._append_chat("全部 PDF 已处理完成。", "success")
            self.root.after(0, lambda: messagebox.showinfo("完成", "PDF 转 JSON / Excel 已完成。"))

        except Exception as e:
            self._append_chat(f"发生错误：{e}", "error")
            self._set_status("失败", badge_bg="#fee2e2", badge_fg="#991b1b")
            self.root.after(0, lambda: messagebox.showerror("错误", str(e)))
        finally:
            self._set_running(False)

    def export_json_to_excel(self, parsed_data: Dict[str, Any], out_path: str):
        wb = Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)

        catalog_ws = wb.create_sheet("目录")
        catalog_ws.append(["页码", "表ID", "表标题", "类型", "工作表名"])
        self._style_header_row(catalog_ws, 1)

        used_sheet_names = {"目录"}

        pages = parsed_data.get("pages", [])
        for page in pages:
            page_no = page.get("page", "")
            tables = page.get("tables", [])

            for table_idx, table in enumerate(tables, start=1):
                table_id = table.get("table_id", f"p{page_no}_t{table_idx}")
                title = table.get("title", "未命名表格")
                table_type = table.get("type", "generic_table")

                base_name = f"P{page_no}_T{table_idx}"
                sheet_name = self._safe_sheet_name(base_name, used_sheet_names)

                ws = wb.create_sheet(sheet_name)
                used_sheet_names.add(sheet_name)

                catalog_ws.append([page_no, table_id, title, table_type, sheet_name])

                self._write_table_block_to_sheet(
                    ws=ws,
                    title=title,
                    headers=table.get("headers", []),
                    data=table.get("data", []),
                    checks=table.get("checks", []),
                    notes=table.get("notes", []),
                    start_row=1
                )

                current_row = ws.max_row + 2

                subtables = table.get("subtables", [])
                for sub_idx, sub in enumerate(subtables, start=1):
                    sub_title = sub.get("title", f"子表{sub_idx}")
                    current_row = self._write_table_block_to_sheet(
                        ws=ws,
                        title=f"[子表{sub_idx}] {sub_title}",
                        headers=sub.get("headers", []),
                        data=sub.get("data", []),
                        checks=sub.get("checks", []),
                        notes=sub.get("notes", []),
                        start_row=current_row
                    )
                    current_row += 2

                self._auto_adjust_sheet(ws)

        self._auto_adjust_sheet(catalog_ws)
        wb.save(out_path)

    def _write_table_block_to_sheet(
        self,
        ws,
        title: str,
        headers: List[List[str]],
        data: List[List[str]],
        checks: List[Dict[str, Any]],
        notes: List[str],
        start_row: int = 1
    ) -> int:
        row_idx = start_row

        ws.cell(row=row_idx, column=1, value=title)
        ws.cell(row=row_idx, column=1).font = Font(bold=True, size=12)
        ws.cell(row=row_idx, column=1).fill = PatternFill("solid", fgColor="DCE6F1")
        ws.cell(row=row_idx, column=1).alignment = Alignment(vertical="center", horizontal="left")
        row_idx += 1

        all_rows = []
        if headers:
            all_rows.extend(headers)
        if data:
            all_rows.extend(data)

        max_cols = max((len(r) for r in all_rows), default=1)

        for r_i, row in enumerate(all_rows, start=row_idx):
            padded = list(row) + [""] * (max_cols - len(row))
            for c_i, value in enumerate(padded, start=1):
                cell = ws.cell(row=r_i, column=c_i, value=value)
                cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)

                if headers and r_i < row_idx + len(headers):
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="EAF2F8")

        row_idx += len(all_rows)

        if checks:
            row_idx += 1
            ws.cell(row=row_idx, column=1, value="Checks")
            ws.cell(row=row_idx, column=1).font = Font(bold=True)
            row_idx += 1

            ws.cell(row=row_idx, column=1, value="Label")
            ws.cell(row=row_idx, column=2, value="Value")
            ws.cell(row=row_idx, column=1).font = Font(bold=True)
            ws.cell(row=row_idx, column=2).font = Font(bold=True)
            row_idx += 1

            for item in checks:
                ws.cell(row=row_idx, column=1, value=item.get("label", ""))
                ws.cell(row=row_idx, column=2, value="是" if item.get("value") else "否")
                row_idx += 1

        if notes:
            row_idx += 1
            ws.cell(row=row_idx, column=1, value="Notes")
            ws.cell(row=row_idx, column=1).font = Font(bold=True)
            row_idx += 1

            for note in notes:
                ws.cell(row=row_idx, column=1, value=note)
                row_idx += 1

        return row_idx

    def _style_header_row(self, ws, row_num: int):
        for cell in ws[row_num]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAD3")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _auto_adjust_sheet(self, ws):
        max_col = ws.max_column
        max_row = ws.max_row

        for col in range(1, max_col + 1):
            max_len = 0
            col_letter = get_column_letter(col)
            for row in range(1, max_row + 1):
                value = ws.cell(row=row, column=col).value
                if value is None:
                    continue
                text = str(value)
                if len(text) > max_len:
                    max_len = len(text)

            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    def _safe_sheet_name(self, name: str, used_names: set) -> str:
        invalid_chars = r'[:\\/?*\[\]]'
        name = re.sub(invalid_chars, "_", name)
        name = name[:31] if len(name) <= 31 else name[:31]
        if not name:
            name = "Sheet"

        original = name
        idx = 1
        while name in used_names:
            suffix = f"_{idx}"
            name = original[:31 - len(suffix)] + suffix
            idx += 1
        return name

    def _append_chat(self, text: str, sender: str = "system"):
        self.root.after(0, lambda: self.chat.add_message(text, sender))

    def _set_status(self, text: str, badge_bg=None, badge_fg=None):
        def update():
            self.status_text.set(text)
            if badge_bg:
                self.status_badge.configure(bg=badge_bg)
            if badge_fg:
                self.status_badge.configure(fg=badge_fg)
        self.root.after(0, update)

    def _set_running(self, running: bool):
        def update():
            if running:
                self.start_btn.state(["disabled"])
                self.json_to_excel_btn.state(["disabled"])
                self.progress.start(10)
            else:
                self.start_btn.state(["!disabled"])
                self.json_to_excel_btn.state(["!disabled"])
                self.progress.stop()
        self.root.after(0, update)


def main():
    root = tk.Tk()
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()