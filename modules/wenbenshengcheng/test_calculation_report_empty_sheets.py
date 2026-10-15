import unittest

from openpyxl import Workbook

from modules.wenbenshengcheng.calculation_report_empty_sheets import (
    collect_component_data_state,
    delete_empty_component_report_sheets,
)


def _make_intermediate(component_rows):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for component_name, rows in component_rows.items():
        sheet = workbook.create_sheet(component_name)
        sheet.append(["结果标识", "参数名称", "参数值"])
        for row in rows:
            sheet.append(row)
    return workbook


def _make_report(sheet_names):
    workbook = Workbook()
    workbook.active.title = sheet_names[0]
    for sheet_name in sheet_names[1:]:
        workbook.create_sheet(sheet_name)
    return workbook


class EmptyComponentReportSheetTests(unittest.TestCase):
    def test_only_final_report_sheets_are_deleted(self):
        intermediate = _make_intermediate({
            "空元件": [("工况1：厚度", "厚度", "")],
            "有数据元件": [("工况1：数量", "数量", 0)],
            "管程入口接管": [("工况1：厚度", "厚度", None)],
            "浮头法兰": [("工况1：厚度", "厚度", "null")],
            "锥壳": [("工况1：厚度", "厚度", "")],
        })
        original_intermediate_sheets = list(intermediate.sheetnames)
        report = _make_report([
            "说明",
            "空元件",
            "有数据元件",
            "管程入口接管校核（内压）",
            "管程入口接管厚壁管小端校核（外压）",
            "浮头法兰（TNC）",
            "内压锥壳",
            "外压锥壳（工况一）",
            "模板辅助页",
        ])

        deleted = delete_empty_component_report_sheets(
            target_sheet_names=list(report.sheetnames),
            intermediate_workbook=intermediate,
            case_label="工况1",
            component_names=[
                "空元件",
                "有数据元件",
                "管程入口接管",
                "浮头法兰",
                "锥壳",
            ],
            delete_sheet=lambda name: report.remove(report[name]),
        )

        self.assertEqual(list(intermediate.sheetnames), original_intermediate_sheets)
        self.assertIn("有数据元件", report.sheetnames)
        self.assertIn("说明", report.sheetnames)
        self.assertIn("模板辅助页", report.sheetnames)
        self.assertNotIn("空元件", report.sheetnames)
        self.assertNotIn("管程入口接管校核（内压）", report.sheetnames)
        self.assertNotIn("管程入口接管厚壁管小端校核（外压）", report.sheetnames)
        self.assertNotIn("浮头法兰（TNC）", report.sheetnames)
        self.assertNotIn("内压锥壳", report.sheetnames)
        self.assertNotIn("外压锥壳（工况一）", report.sheetnames)
        self.assertEqual(
            deleted,
            {
                "空元件",
                "管程入口接管校核（内压）",
                "管程入口接管厚壁管小端校核（外压）",
                "浮头法兰（TNC）",
                "内压锥壳",
                "外压锥壳（工况一）",
            },
        )

    def test_case_data_is_judged_independently(self):
        intermediate = _make_intermediate({
            "多工况元件": [
                ("工况1：厚度", "厚度", 12.5),
                ("工况2：厚度", "厚度", ""),
            ],
        })

        empty_case_1, nonempty_case_1 = collect_component_data_state(
            intermediate, "工况一", ["多工况元件"]
        )
        empty_case_2, nonempty_case_2 = collect_component_data_state(
            intermediate, "工况2", ["多工况元件"]
        )

        self.assertEqual(empty_case_1, set())
        self.assertEqual(nonempty_case_1, {"多工况元件"})
        self.assertEqual(empty_case_2, {"多工况元件"})
        self.assertEqual(nonempty_case_2, set())

    def test_existing_cone_logic_remains_authoritative_when_any_cone_has_data(self):
        intermediate = _make_intermediate({
            "锥壳": [("工况1：厚度", "厚度", "")],
            "锥壳(α1-工况一)": [("工况1：厚度", "厚度", 8.0)],
        })
        report = _make_report(["说明", "内压锥壳", "外压锥壳（工况一）"])

        deleted = delete_empty_component_report_sheets(
            target_sheet_names=list(report.sheetnames),
            intermediate_workbook=intermediate,
            case_label="工况1",
            component_names=["锥壳", "锥壳(α1-工况一)"],
            delete_sheet=lambda name: report.remove(report[name]),
        )

        self.assertEqual(deleted, set())
        self.assertIn("内压锥壳", report.sheetnames)
        self.assertIn("外压锥壳（工况一）", report.sheetnames)


if __name__ == "__main__":
    unittest.main()
