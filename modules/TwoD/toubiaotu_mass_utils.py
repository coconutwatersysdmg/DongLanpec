import json
import os
import re
import time
import traceback

import openpyxl


def _extract_number(value):
    """从单元格值中提取数字；公式字符串不强行解析，避免把 =SUM(J8:J20) 误读成 8。"""
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    s = str(value).strip()
    if s == "" or s.lower() in ("none", "null", "nan"):
        return None

    if s.startswith("="):
        return None

    s = s.replace(",", "").replace("，", "")

    m = re.search(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?", s)
    return float(m.group()) if m else None


def _json_safe(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _cell_value(ws, cell_name):
    value = ws[cell_name].value
    if value is not None:
        return value

    for merged_range in ws.merged_cells.ranges:
        if cell_name in merged_range:
            return ws.cell(
                row=merged_range.min_row,
                column=merged_range.min_col
            ).value

    return None


def _is_blank(value):
    if value is None:
        return True
    return str(value).strip() in ("", "None", "none", "NULL", "null")


def _display(value):
    if value is None:
        return ""
    return str(value)


def _format_audit_text(audit):
    lines = [
        "========== 7785F 质量审计日志 ==========",
        f"生成时间: {audit.get('generated_at')}",
        f"来源文件: {audit.get('excel_path')}",
        f"工作表: {audit.get('sheet_name')}",
        f"模式: {audit.get('mode')}",
        f"调用文件: {audit.get('audit_label') or ''}",
    ]

    if audit.get("workbook_sheets"):
        lines.append(f"工作簿包含工作表: {', '.join(audit.get('workbook_sheets') or [])}")

    result = audit.get("result") or {}
    if result:
        lines.extend([
            "",
            "[计算结果]",
            f"Sheet1采用总质量: {result.get('sheet_total')}，来源: {result.get('sheet_total_source')}",
            f"U形换热管重量: {result.get('tube_total')}",
            f"7785F函数返回值: {result.get('final')}",
        ])
        if "expected_written_value_after_existing_div2" in result:
            lines.append(f"当前调用点 /2 后预计写图值: {result.get('expected_written_value_after_existing_div2')}")

    warnings = audit.get("warnings") or []
    if warnings:
        lines.append("")
        lines.append("[警告]")
        lines.extend(f"- {warning}" for warning in warnings)

    u_tube = audit.get("u_tube") or {}
    if u_tube:
        lines.extend([
            "",
            "[U形换热管来源]",
            f"采用来源: {u_tube.get('source')}",
            f"U形管明细!G6 原始值: {repr(u_tube.get('g6_raw'))}",
            f"U形管明细!G6 数值: {u_tube.get('g6_num')}",
            f"最终U形换热管重量: {u_tube.get('tube_total')}",
        ])
        sheet1_u_row = u_tube.get("sheet1_u_tube_row") or {}
        if sheet1_u_row:
            lines.append(
                "Sheet1 U形换热管行: "
                f"row={sheet1_u_row.get('row')}, "
                f"I原始={repr(sheet1_u_row.get('i_raw'))}, I数值={sheet1_u_row.get('i_num')}, "
                f"J原始={repr(sheet1_u_row.get('j_raw'))}, J数值={sheet1_u_row.get('j_num')}"
            )

    sheet1 = audit.get("sheet1") or {}
    if sheet1:
        lines.extend([
            "",
            "[Sheet1 汇总]",
            f"明细I列合计: {sheet1.get('data_i_sum')}",
            f"明细J列合计: {sheet1.get('data_j_sum')}",
            f"所有J列合计(含总计行): {sheet1.get('all_j_sum')}",
            f"识别到的总计行J列: {sheet1.get('total_row_j')}",
            f"最后一个有效J列: {sheet1.get('last_j')}",
        ])

        rows = sheet1.get("rows") or []
        lines.extend([
            "",
            "[Sheet1 明细行 D/G/H/I/J]",
            "行号 | 元件名称 | G列原始 | H列原始 | I列原始 -> I数值 | J列原始 -> J数值 | 标记",
        ])
        for row in rows:
            marks = []
            if row.get("is_total"):
                marks.append("总计行")
            if row.get("name") and row.get("i_num") is None and row.get("j_num") is None and not row.get("is_total"):
                marks.append("I/J缺有效数字")
            lines.append(
                f"{row.get('row')} | "
                f"{_display(row.get('name'))} | "
                f"{repr(row.get('qty_g_raw'))} | "
                f"{repr(row.get('h_raw'))} | "
                f"{repr(row.get('i_raw'))} -> {row.get('i_num')} | "
                f"{repr(row.get('j_raw'))} -> {row.get('j_num')} | "
                f"{'，'.join(marks)}"
            )

    if audit.get("error"):
        lines.extend(["", "[异常]", audit.get("error")])

    lines.append("========== 7785F 质量审计日志结束 ==========")
    return "\n".join(lines)


def _print_audit_log(audit):
    try:
        print(_format_audit_text(audit))
    except Exception:
        print(f"⚠️ 7785F质量审计日志打印失败:\n{traceback.format_exc()}")


def _write_audit_file(excel_path, audit):
    try:
        base_dir = os.path.dirname(os.path.abspath(excel_path)) or os.getcwd()
        audit_dir = os.path.join(base_dir, "7785F_audit")
        os.makedirs(audit_dir, exist_ok=True)

        timestamp = time.strftime("%Y%m%d_%H%M%S")
        unique_suffix = time.time_ns()
        pid = os.getpid()
        mode = re.sub(r"[^0-9A-Za-z_\-]+", "_", str(audit.get("mode") or "unknown"))
        audit_path = os.path.join(audit_dir, f"7785F_audit_{timestamp}_{pid}_{mode}_{unique_suffix}.json")
        txt_path = os.path.splitext(audit_path)[0] + ".txt"

        with open(audit_path, "w", encoding="utf-8") as f:
            json.dump(audit, f, ensure_ascii=False, indent=2)

        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(_format_audit_text(audit))

        print(f"🧾 7785F质量审计JSON: {audit_path}")
        print(f"🧾 7785F质量审计TXT: {txt_path}")
        return audit_path, txt_path
    except Exception:
        print(f"⚠️ 7785F质量审计文件写入失败:\n{traceback.format_exc()}")
        return None


def _collect_sheet1_mass(ws, min_row=8):
    rows = []
    data_j_sum = 0.0
    data_i_sum = 0.0
    all_j_sum = 0.0
    total_rows = []
    last_j = None
    warnings = []
    missing_rows = []
    u_tube_row = None

    for r in range(min_row, ws.max_row + 1):
        name_raw = ws.cell(row=r, column=4).value   # D列：元件名称
        qty_raw = ws.cell(row=r, column=7).value    # G列：数量
        h_raw = ws.cell(row=r, column=8).value      # H列：单重/总计标识
        i_raw = ws.cell(row=r, column=9).value      # I列：质量
        j_raw = ws.cell(row=r, column=10).value     # J列：成品质量/总计

        if all(_is_blank(v) for v in (name_raw, qty_raw, h_raw, i_raw, j_raw)):
            continue

        name = str(name_raw).strip() if not _is_blank(name_raw) else ""
        h_text = str(h_raw).strip() if not _is_blank(h_raw) else ""
        is_total = name == "总计" or h_text == "总计"

        i_num = _extract_number(i_raw)
        j_num = _extract_number(j_raw)

        row_info = {
            "row": r,
            "name": name,
            "qty_g_raw": _json_safe(qty_raw),
            "h_raw": _json_safe(h_raw),
            "i_raw": _json_safe(i_raw),
            "j_raw": _json_safe(j_raw),
            "i_num": i_num,
            "j_num": j_num,
            "is_total": is_total,
        }
        rows.append(row_info)

        if j_num is not None:
            all_j_sum += j_num
            last_j = {
                "row": r,
                "name": name,
                "h_raw": _json_safe(h_raw),
                "j_raw": _json_safe(j_raw),
                "j_num": j_num,
                "is_total": is_total,
            }

        if is_total:
            if j_num is not None:
                total_rows.append(row_info)
            continue

        if name:
            if i_num is not None:
                data_i_sum += i_num
            if j_num is not None:
                data_j_sum += j_num

            if name == "U形换热管":
                u_tube_row = row_info

            if i_num is None and j_num is None:
                missing_rows.append(row_info)

    total_row_j = None
    if total_rows:
        total_row_j = total_rows[-1]["j_num"]

    if missing_rows:
        warnings.append(
            "Sheet1存在元件行但I/J列都不是有效数字: "
            + "，".join(f"{row['row']}:{row['name']}" for row in missing_rows[:20])
        )

    return {
        "rows": rows,
        "data_i_sum": round(data_i_sum, 6),
        "data_j_sum": round(data_j_sum, 6),
        "all_j_sum": round(all_j_sum, 6),
        "total_rows": total_rows,
        "total_row_j": total_row_j,
        "last_j": last_j,
        "missing_rows": missing_rows,
        "u_tube_row": u_tube_row,
        "warnings": warnings,
    }


def _choose_sheet_total(sheet_audit, warnings):
    total_row_j = sheet_audit.get("total_row_j")
    data_j_sum = sheet_audit.get("data_j_sum") or 0.0
    last_j = sheet_audit.get("last_j")

    if total_row_j is None:
        if last_j is None:
            warnings.append("Sheet1未找到J列总计，也没有任何有效J列数字，Sheet1总质量按0处理")
            return 0.0, "no_j_value"
        warnings.append("Sheet1未找到明确的“总计”行，改用明细J列求和作为总质量")
        return data_j_sum, "sum_detail_j_no_total_row"

    if data_j_sum > total_row_j + 0.01:
        warnings.append(
            f"Sheet1明细J列求和({data_j_sum})大于总计行J列({total_row_j})，改用明细J列求和"
        )
        return data_j_sum, "sum_detail_j_gt_total_row"

    return float(total_row_j), "total_row_j"


def _collect_utube_mass(wb, sheet_audit, warnings):
    g6_raw = None
    g6_num = None

    if "U形管明细" not in wb.sheetnames:
        warnings.append("未找到工作表“U形管明细”，U形换热管重量将尝试从Sheet1兜底")
    else:
        ws_u = wb["U形管明细"]
        g6_raw = _cell_value(ws_u, "G6")
        g6_num = _extract_number(g6_raw)
        if g6_num is None:
            warnings.append(f"U形管明细!G6不是有效数字: {repr(g6_raw)}")

    sheet_u_row = sheet_audit.get("u_tube_row")
    sheet_u_i = sheet_u_row.get("i_num") if sheet_u_row else None

    if g6_num is not None:
        source = "U形管明细!G6"
        tube_total = g6_num
    elif sheet_u_i is not None:
        source = "Sheet1 U形换热管 I列兜底"
        tube_total = sheet_u_i
        warnings.append("U形管明细!G6未读到有效数字，已使用Sheet1中U形换热管I列兜底")
    else:
        source = "missing"
        tube_total = 0.0
        warnings.append("U形换热管重量未读到，按0处理")

    if g6_num is not None and sheet_u_i is not None and abs(g6_num - sheet_u_i) > 0.01:
        warnings.append(
            f"U形管明细!G6({g6_num}) 与 Sheet1 U形换热管I列({sheet_u_i}) 不一致"
        )

    return {
        "g6_raw": _json_safe(g6_raw),
        "g6_num": g6_num,
        "sheet1_u_tube_row": sheet_u_row,
        "source": source,
        "tube_total": tube_total,
    }


def calculate_7785f_from_excel_with_audit(excel_path, sheet_name, mode="sheet_total_plus_utube", audit_label=None):
    """
    统一计算投标图 7785F 质量，并输出审计 JSON。

    mode:
      - sheet_total_plus_utube: U型管图纸，结果 = Sheet1总计J + U形管明细G6
      - double_sheet_total_for_half_write: 保持旧调用点“返回值/2”写图的行为，结果 = 2 * 可靠Sheet1总计
    """
    excel_path = os.path.abspath(excel_path)
    warnings = []
    audit = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "excel_path": excel_path,
        "sheet_name": sheet_name,
        "mode": mode,
        "audit_label": audit_label,
        "warnings": warnings,
    }

    print(f"📌 开始计算7785F，读取文件: {excel_path}，mode={mode}")

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        audit["workbook_sheets"] = list(wb.sheetnames)

        if sheet_name not in wb.sheetnames:
            warnings.append(f"未找到工作表“{sheet_name}”，7785F按0处理")
            audit["result"] = {
                "sheet_total": 0.0,
                "sheet_total_source": "missing_sheet",
                "tube_total": 0.0,
                "final": 0.0,
            }
            return 0.0

        ws = wb[sheet_name]
        sheet_audit = _collect_sheet1_mass(ws)
        audit["sheet1"] = sheet_audit
        warnings.extend(sheet_audit.pop("warnings", []))

        sheet_total, sheet_total_source = _choose_sheet_total(sheet_audit, warnings)
        tube_total = 0.0

        if mode == "sheet_total_plus_utube":
            u_audit = _collect_utube_mass(wb, sheet_audit, warnings)
            audit["u_tube"] = u_audit
            tube_total = u_audit["tube_total"]

            u_row = sheet_audit.get("u_tube_row")
            if u_row and u_row.get("j_num") and abs(float(u_row["j_num"])) > 0.01:
                warnings.append(
                    f"Sheet1中U形换热管J列已有数值({u_row['j_num']})，请确认是否与G6重复计入"
                )

            final = sheet_total + tube_total
        elif mode == "double_sheet_total_for_half_write":
            final = sheet_total * 2
        else:
            warnings.append(f"未知7785F计算模式: {mode}，使用Sheet1总计J列")
            final = sheet_total

        audit["result"] = {
            "sheet_total": round(sheet_total, 6),
            "sheet_total_source": sheet_total_source,
            "tube_total": round(tube_total, 6),
            "final": round(final, 6),
        }
        if mode == "double_sheet_total_for_half_write":
            audit["result"]["expected_written_value_after_existing_div2"] = round(final / 2.0, 6)

        _print_audit_log(audit)
        audit["_printed"] = True

        return final
    except Exception:
        err = traceback.format_exc()
        audit["error"] = err
        warnings.append("calculate_7785f_from_excel_with_audit 出错")
        print(f"❌ calculate_7785f_from_excel_with_audit 出错:\n{err}")
        return 0.0
    finally:
        try:
            if "wb" in locals():
                wb.close()
        finally:
            if not audit.pop("_printed", False):
                _print_audit_log(audit)
            _write_audit_file(excel_path, audit)
